# Prints every tab of an .xlsx as the two grids the Sheets API hands the importer:
# UNFORMATTED_VALUE (cached values) and FORMULA (formula text, else the literal).
import json
import sys

import openpyxl

LAST_COLUMN_INDEX = 78  # BZ, the importer's read range


def grid(sheet):
    rows = []
    for row in sheet.iter_rows(max_col=min(sheet.max_column, LAST_COLUMN_INDEX)):
        cells = ["" if cell.value is None else cell.value for cell in row]
        while cells and cells[-1] == "":
            cells.pop()
        rows.append(cells)
    while rows and not rows[-1]:
        rows.pop()
    return rows


path = sys.argv[1]
formulas = openpyxl.load_workbook(path)
values = openpyxl.load_workbook(path, data_only=True)
tabs = [
    {"title": sheet.title, "values": grid(values[sheet.title]), "formulas": grid(sheet)}
    for sheet in formulas.worksheets
]
json.dump(tabs, sys.stdout, default=str)
